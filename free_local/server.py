from __future__ import annotations

import json
import re
import socket
import threading
import webbrowser
from pathlib import Path
from datetime import datetime
from typing import Any

import requests
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from faster_whisper import WhisperModel

APP_DIR = Path(__file__).resolve().parent
STATIC_DIR = APP_DIR / "static"
CONFIG_PATH = APP_DIR / "config.json"

DEFAULT_CONFIG = {
    "drive_root": "",
    "whisper_model": "small",
    "ollama_model": "qwen2.5:3b",
}

app = FastAPI(title="Gestor de Informes FM Ciudad - Gratis")
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

_whisper: WhisperModel | None = None
_whisper_name: str | None = None


def load_config() -> dict[str, Any]:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        try:
            cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
        except Exception:
            pass
    return cfg


def save_config(cfg: dict[str, Any]) -> None:
    CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


def safe_title(value: str) -> str:
    text = re.sub(r'[\\/:*?"<>|]+', "", value or "")
    text = re.sub(r"\s+", " ", text).strip()
    return (text[:90] or "Sin título").strip()


def safe_ext(filename: str, fallback: str) -> str:
    suffix = Path(filename or "").suffix.lower().lstrip(".")
    return suffix or fallback


def get_whisper(model_name: str) -> WhisperModel:
    global _whisper, _whisper_name
    if _whisper is None or _whisper_name != model_name:
        _whisper = WhisperModel(model_name, device="cpu", compute_type="int8")
        _whisper_name = model_name
    return _whisper


def transcribe(audio_path: Path, model_name: str) -> str:
    model = get_whisper(model_name)
    segments, _ = model.transcribe(
        str(audio_path),
        language="es",
        vad_filter=True,
        beam_size=5,
    )
    text = " ".join(seg.text.strip() for seg in segments if seg.text.strip())
    return re.sub(r"\s+", " ", text).strip()


def fallback_analysis(transcript: str) -> dict[str, str]:
    clean = re.sub(r"\s+", " ", transcript).strip()
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", clean) if s.strip()]
    title_seed = sentences[0] if sentences else clean
    title_words = title_seed.split()[:12]
    title = safe_title(" ".join(title_words))
    short = " ".join(sentences[:2]) if sentences else clean[:280]
    expanded = " ".join(sentences[:5]) if sentences else clean[:900]
    words = expanded.split()
    if len(words) > 120:
        expanded = " ".join(words[:120]).rstrip(",;:") + "."
    return {
        "title": title or "Informe",
        "short_summary": short,
        "expanded_summary": expanded,
    }


def analyze_with_ollama(transcript: str, model_name: str) -> tuple[dict[str, str], bool]:
    prompt = f"""
Sos editor periodístico de FM Ciudad 89.7 de Río Gallegos.
Trabajá SOLO con la transcripción que te doy.

Generá:
1. un título periodístico breve y claro;
2. un resumen corto de 1 o 2 oraciones;
3. un resumen ampliado de UN solo párrafo, entre 70 y 120 palabras.

No inventes datos, nombres, cargos, lugares ni cifras.
No agregues opinión.
Respondé únicamente JSON válido con estas claves exactas:
{{"title":"...","short_summary":"...","expanded_summary":"..."}}

TRANSCRIPCIÓN:
{transcript}
""".strip()
    try:
        response = requests.post(
            "http://127.0.0.1:11434/api/generate",
            json={
                "model": model_name,
                "prompt": prompt,
                "stream": False,
                "format": "json",
                "options": {"temperature": 0.2},
            },
            timeout=360,
        )
        response.raise_for_status()
        payload = response.json()
        data = json.loads(payload.get("response", "{}"))
        result = {
            "title": safe_title(str(data.get("title", ""))),
            "short_summary": str(data.get("short_summary", "")).strip(),
            "expanded_summary": str(data.get("expanded_summary", "")).strip(),
        }
        if not result["title"] or not result["expanded_summary"]:
            raise ValueError("Respuesta incompleta de Ollama")
        return result, True
    except Exception:
        return fallback_analysis(transcript), False


def create_word(date_label: str, reports: list[dict[str, Any]], destination: Path) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Resumen de informes - FM Ciudad 89.7", level=0)
    title.alignment = 1
    p = doc.add_paragraph()
    run = p.add_run(f"Fecha: {date_label}")
    run.bold = True

    for report in reports:
        number = str(report["number"]).zfill(2)
        doc.add_heading(f"Informe {number} - {report['title']}", level=1)
        doc.add_paragraph(report["expanded_summary"])

    doc.save(str(destination))


def create_txt(reports: list[dict[str, Any]], destination: Path) -> None:
    chunks = []
    for report in reports:
        number = str(report["number"]).zfill(2)
        chunks.append(
            f"Informe {number} - {report['title']}\n{report['expanded_summary']}"
        )
    destination.write_text("\n\n".join(chunks), encoding="utf-8")


def date_folder_name(date_iso: str) -> str:
    dt = datetime.strptime(date_iso, "%Y-%m-%d")
    return dt.strftime("%d-%m-%Y")


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/status")
def status():
    cfg = load_config()
    root = Path(cfg["drive_root"]) if cfg.get("drive_root") else None
    return {
        "ok": True,
        "drive_root": str(root) if root else "",
        "drive_root_exists": bool(root and root.exists()),
        "whisper_model": cfg["whisper_model"],
        "ollama_model": cfg["ollama_model"],
        "ollama_online": _ollama_online(),
        "lan_url": f"http://{_local_ip()}:8765",
    }


@app.post("/api/settings")
async def settings(request: Request):
    data = await request.json()
    cfg = load_config()

    root = str(data.get("drive_root", "")).strip()
    if root:
        path = Path(root).expanduser()
        try:
            path.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            return JSONResponse({"error": f"No se puede usar esa carpeta: {exc}"}, status_code=400)
        cfg["drive_root"] = str(path.resolve())

    if data.get("whisper_model"):
        cfg["whisper_model"] = str(data["whisper_model"]).strip()
    if data.get("ollama_model"):
        cfg["ollama_model"] = str(data["ollama_model"]).strip()

    save_config(cfg)
    return {"ok": True, **cfg}


@app.post("/api/select-folder")
def select_folder():
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        selected = filedialog.askdirectory(title="Elegí la carpeta raíz sincronizada con Google Drive")
        root.destroy()
        if not selected:
            return {"ok": False, "cancelled": True}

        cfg = load_config()
        cfg["drive_root"] = str(Path(selected).resolve())
        save_config(cfg)
        return {"ok": True, "drive_root": cfg["drive_root"]}
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=500)


@app.post("/api/process-batch")
async def process_batch(request: Request):
    cfg = load_config()
    root_text = cfg.get("drive_root", "").strip()
    if not root_text:
        return JSONResponse({"error": "Primero elegí la carpeta raíz de Google Drive."}, status_code=400)

    root = Path(root_text)
    root.mkdir(parents=True, exist_ok=True)

    form = await request.form()
    date_iso = str(form.get("date", "")).strip()
    try:
        folder_name = date_folder_name(date_iso)
    except Exception:
        return JSONResponse({"error": "Fecha inválida."}, status_code=400)

    daily = root / folder_name
    daily.mkdir(parents=True, exist_ok=True)

    indexes = sorted(
        {
            int(match.group(1))
            for key in form.keys()
            if (match := re.fullmatch(r"audio_(\d+)", key))
        }
    )
    if not indexes:
        return JSONResponse({"error": "No se recibieron informes."}, status_code=400)

    reports: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []

    for idx in indexes:
        audio = form.get(f"audio_{idx}")
        photo = form.get(f"photo_{idx}")
        if not audio or not photo:
            failures.append({"number": idx, "error": "Falta audio o foto."})
            continue

        number = str(idx).zfill(2)
        audio_ext = safe_ext(getattr(audio, "filename", ""), "mp3")
        photo_ext = safe_ext(getattr(photo, "filename", ""), "jpg")
        temp_audio = daily / f"Informe {number} - procesando.{audio_ext}"
        temp_photo = daily / f"Informe {number} - procesando.{photo_ext}"

        try:
            temp_audio.write_bytes(await audio.read())
            temp_photo.write_bytes(await photo.read())

            transcript = transcribe(temp_audio, cfg["whisper_model"])
            if not transcript:
                raise RuntimeError("No se pudo detectar voz en el audio.")

            analysis, used_ollama = analyze_with_ollama(transcript, cfg["ollama_model"])
            title = safe_title(analysis["title"])

            final_audio = daily / f"Informe {number} - {title}.{audio_ext}"
            final_photo = daily / f"Informe {number} - {title}.{photo_ext}"

            if final_audio.exists():
                final_audio.unlink()
            if final_photo.exists():
                final_photo.unlink()

            temp_audio.rename(final_audio)
            temp_photo.rename(final_photo)

            reports.append(
                {
                    "number": idx,
                    "title": title,
                    "short_summary": analysis["short_summary"],
                    "expanded_summary": analysis["expanded_summary"],
                    "transcript": transcript,
                    "audio_file": final_audio.name,
                    "photo_file": final_photo.name,
                    "used_ollama": used_ollama,
                }
            )
        except Exception as exc:
            failures.append({"number": idx, "error": str(exc)})

    reports.sort(key=lambda item: item["number"])

    if not reports:
        return JSONResponse(
            {"error": "No se pudo procesar ningún informe.", "failures": failures},
            status_code=500,
        )

    docx_path = daily / f"Resumen de informes {folder_name}.docx"
    txt_path = daily / f"Resumen de informes {folder_name}.txt"
    create_word(folder_name, reports, docx_path)
    create_txt(reports, txt_path)

    return {
        "ok": True,
        "folder": str(daily),
        "date_folder": folder_name,
        "reports": reports,
        "failures": failures,
        "word_file": docx_path.name,
        "text_file": txt_path.name,
        "ollama_online": _ollama_online(),
    }


def _ollama_online() -> bool:
    try:
        r = requests.get("http://127.0.0.1:11434/api/tags", timeout=1.5)
        return r.ok
    except Exception:
        return False


def _local_ip() -> str:
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
        return ip
    except Exception:
        return "127.0.0.1"


def open_browser_later() -> None:
    def _open():
        import time
        time.sleep(1.2)
        webbrowser.open("http://127.0.0.1:8765")
    threading.Thread(target=_open, daemon=True).start()


if __name__ == "__main__":
    import uvicorn

    open_browser_later()
    print("")
    print("Gestor de Informes FM Ciudad - modo GRATIS")
    print("PC:   http://127.0.0.1:8765")
    print(f"WiFi: http://{_local_ip()}:8765")
    print("")
    uvicorn.run(app, host="0.0.0.0", port=8765)
